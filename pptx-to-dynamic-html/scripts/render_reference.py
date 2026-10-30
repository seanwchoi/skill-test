#!/usr/bin/env python3
"""Render reference images and crop raster fallbacks.

Usage:
    python render_reference.py deck.pptx OUT_DIR      (after extract_pptx.py)

1. OUT_DIR/ref/slide-N.png  - the original deck rendered by LibreOffice. Used to
   compare the HTML against the original (design fidelity check).
2. For every element extracted as kind="raster" (charts, tables, SmartArt,
   unusual geometry), renders a copy of the deck in which only those elements
   remain on each slide, crops each one out, saves OUT_DIR/media/raster-*.png and
   writes the path into slides.json as "src". Rendering a stripped copy keeps
   callouts and other shapes from being baked into the crops.
"""
import glob
import json
import os
import shutil
import subprocess
import sys
import tempfile

from lxml import etree
from PIL import Image
from pptx import Presentation

RENDER_W = 2400  # px width of reference renders


def to_pngs(pptx_path, out_dir, prefix):
    os.makedirs(out_dir, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", td, pptx_path],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
        pdf = glob.glob(os.path.join(td, "*.pdf"))[0]
        subprocess.run(["pdftoppm", "-png", "-scale-to-x", str(RENDER_W), "-scale-to-y", "-1",
                        pdf, os.path.join(td, "p")], check=True)
        files = sorted(glob.glob(os.path.join(td, "p-*.png")),
                       key=lambda f: int(f.rsplit("-", 1)[1].split(".")[0]))
        out = []
        for i, f in enumerate(files, 1):
            dst = os.path.join(out_dir, "%s-%d.png" % (prefix, i))
            shutil.move(f, dst)
            out.append(dst)
        return out


def main(pptx_path, out_dir):
    data_path = os.path.join(out_dir, "slides.json")
    data = json.load(open(data_path, encoding="utf-8"))
    to_pngs(pptx_path, os.path.join(out_dir, "ref"), "slide")

    keep = {s["index"]: {e["id"] for e in s["elements"] if e["kind"] == "raster" and e["origin"] == "slide"}
            for s in data["slides"]}
    any_raster = any(e["kind"] == "raster" for s in data["slides"] for e in s["elements"])
    if not any_raster:
        print("no raster elements; reference renders in", os.path.join(out_dir, "ref"))
        return

    prs = Presentation(pptx_path)
    for i, slide in enumerate(prs.slides, 1):
        tree = slide.shapes._spTree
        for child in list(tree):
            tag = etree.QName(child).localname
            if tag in ("nvGrpSpPr", "grpSpPr", "extLst"):
                continue
            ids = {c.get("id") for c in child.iter() if etree.QName(c).localname == "cNvPr"}
            if not ids & keep.get(i, set()):
                tree.remove(child)
        timing = slide.element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}timing")
        if timing is not None:
            slide.element.remove(timing)
    with tempfile.TemporaryDirectory() as td:
        stripped = os.path.join(td, "raster_only.pptx")
        prs.save(stripped)
        pngs = to_pngs(stripped, os.path.join(td, "r"), "r")
        sw, sh = data["width"], data["height"]
        for s, png in zip(data["slides"], pngs):
            img = Image.open(png)
            k = img.width / sw
            for e in s["elements"]:
                if e["kind"] != "raster":
                    continue
                box = (max(0, int(e["x"] * k)), max(0, int(e["y"] * k)),
                       min(img.width, int((e["x"] + e["w"]) * k + 0.999)),
                       min(img.height, int((e["y"] + e["h"]) * k + 0.999)))
                if box[2] <= box[0] or box[3] <= box[1]:
                    continue
                name = "media/raster-s%d-%s.png" % (s["index"], e["id"])
                img.crop(box).save(os.path.join(out_dir, name))
                e["src"] = name
    json.dump(data, open(data_path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("cropped raster elements into", os.path.join(out_dir, "media"))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])
