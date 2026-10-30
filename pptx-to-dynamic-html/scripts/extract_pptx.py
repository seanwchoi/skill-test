#!/usr/bin/env python3
"""Extract a .pptx into a slides.json the HTML builder can consume.

Usage:
    python extract_pptx.py deck.pptx OUT_DIR

Writes OUT_DIR/slides.json and OUT_DIR/media/* (embedded pictures).
Everything is measured in points (1pt = 12700 EMU) in slide coordinates,
with theme/scheme colours already resolved to #RRGGBB(+alpha).

What gets extracted per element ("kind"):
  shape      preset geometry (prst) or custom geometry (path list), fill, line, text
  picture    image file + crop
  line       straight / bent / curved connectors and plain lines, with arrow ends
  raster     anything we do not redraw faithfully (charts, tables, SmartArt,
             unknown geometry, OLE, media) -> builder crops it from a render
Plus: callout detection (+ tail tip point), slide background, entrance-animation
sequence from <p:timing>, and master/layout decoration shapes.
"""
import colorsys
import copy
import json
import os
import re
import sys

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
A = "{%s}" % NS["a"]
P = "{%s}" % NS["p"]
R = "{%s}" % NS["r"]
EMU_PT = 12700.0

# Geometries the builder knows how to draw as SVG. Anything else -> raster.
SUPPORTED_PRST = {
    "rect", "roundRect", "snip1Rect", "snip2SameRect", "round1Rect", "round2SameRect",
    "ellipse", "triangle", "rtTriangle", "diamond", "parallelogram", "trapezoid",
    "pentagon", "homePlate", "chevron", "hexagon", "octagon", "plus", "donut",
    "rightArrow", "leftArrow", "upArrow", "downArrow", "leftRightArrow", "upDownArrow",
    "notchedRightArrow", "stripedRightArrow",
    "wedgeRectCallout", "wedgeRoundRectCallout", "wedgeEllipseCallout", "cloudCallout",
    "flowChartProcess", "flowChartAlternateProcess", "flowChartDecision",
    "flowChartTerminator", "flowChartConnector",
    "line", "straightConnector1",
    "bentConnector2", "bentConnector3", "bentConnector4",
    "curvedConnector2", "curvedConnector3", "curvedConnector4",
}
LINE_PRST = {"line", "straightConnector1", "bentConnector2", "bentConnector3",
             "bentConnector4", "bentConnector5", "curvedConnector2", "curvedConnector3",
             "curvedConnector4", "curvedConnector5"}
CALLOUT_PRST = {"wedgeRectCallout", "wedgeRoundRectCallout", "wedgeEllipseCallout",
                "cloudCallout", "borderCallout1", "borderCallout2", "borderCallout3",
                "callout1", "callout2", "callout3", "accentCallout1", "accentCallout2",
                "accentCallout3", "accentBorderCallout1", "accentBorderCallout2",
                "accentBorderCallout3"}
CALLOUT_NAME_RE = re.compile(r"(말풍선|callout|tooltip|speech\s*bubble|balloon)", re.I)
ARROW_PRST = {"rightArrow", "leftArrow", "upArrow", "downArrow", "leftRightArrow",
              "upDownArrow", "notchedRightArrow", "stripedRightArrow", "chevron",
              "homePlate", "pentagon"}

ENTRANCE = {  # presetID -> effect name used by the builder
    1: "appear", 2: "fly", 3: "fade", 9: "fade", 10: "fade", 12: "fly-short",
    16: "split", 17: "zoom", 18: "wipe", 21: "zoom", 22: "wipe", 23: "zoom",
    26: "zoom", 31: "zoom", 37: "fly-short", 42: "fly-short", 47: "fly-short",
    53: "zoom", 55: "fade", 47: "fly-short",
}
DIRECTION = {1: "top", 2: "right", 4: "bottom", 8: "left", 3: "right", 6: "bottom",
             9: "left", 12: "bottom", 0: "bottom", 10: "horizontal", 5: "vertical",
             21: "horizontal", 26: "vertical", 37: "horizontal", 42: "vertical"}


# --------------------------------------------------------------------------- colours
def _hex_to_rgb(h):
    h = h.lstrip("#")
    return tuple(int(h[i:i + 2], 16) / 255.0 for i in (0, 2, 4))


def _rgb_to_hex(rgb):
    return "#%02X%02X%02X" % tuple(max(0, min(255, round(c * 255))) for c in rgb)


SYS_DEFAULT = {"windowText": "000000", "window": "FFFFFF"}
PRESET_COLORS = {"black": "000000", "white": "FFFFFF", "red": "FF0000",
                 "green": "008000", "blue": "0000FF", "yellow": "FFFF00",
                 "gray": "808080", "grey": "808080", "orange": "FFA500"}


class Theme:
    def __init__(self, theme_el, clr_map):
        self.colors = {}
        self.major = "Calibri Light"
        self.minor = "Calibri"
        self.major_ea = ""
        self.minor_ea = ""
        self.ln_widths = [6350, 12700, 19050]
        self.fill_styles = []
        self.bg_styles = []
        if theme_el is not None:
            cs = theme_el.find(".//a:clrScheme", NS)
            if cs is not None:
                for child in cs:
                    tag = etree.QName(child).localname
                    c = child[0]
                    if etree.QName(c).localname == "srgbClr":
                        self.colors[tag] = c.get("val")
                    elif etree.QName(c).localname == "sysClr":
                        self.colors[tag] = c.get("lastClr") or SYS_DEFAULT.get(c.get("val"), "000000")
            fs = theme_el.find(".//a:fontScheme", NS)
            if fs is not None:
                mj = fs.find("a:majorFont", NS)
                mn = fs.find("a:minorFont", NS)
                if mj is not None:
                    self.major = mj.find("a:latin", NS).get("typeface")
                    ea = mj.find("a:ea", NS)
                    self.major_ea = ea.get("typeface") if ea is not None else ""
                    hang = mj.find("a:font[@script='Hang']", NS)
                    if hang is not None and not self.major_ea:
                        self.major_ea = hang.get("typeface")
                if mn is not None:
                    self.minor = mn.find("a:latin", NS).get("typeface")
                    ea = mn.find("a:ea", NS)
                    self.minor_ea = ea.get("typeface") if ea is not None else ""
                    hang = mn.find("a:font[@script='Hang']", NS)
                    if hang is not None and not self.minor_ea:
                        self.minor_ea = hang.get("typeface")
            lsl = theme_el.find(".//a:lnStyleLst", NS)
            if lsl is not None:
                self.ln_widths = [int(ln.get("w", 12700)) for ln in lsl.findall("a:ln", NS)]
            fsl = theme_el.find(".//a:fillStyleLst", NS)
            if fsl is not None:
                self.fill_styles = list(fsl)
            bsl = theme_el.find(".//a:bgFillStyleLst", NS)
            if bsl is not None:
                self.bg_styles = list(bsl)
        self.clr_map = clr_map or {}

    def scheme(self, name, ph_color=None):
        if name == "phClr" and ph_color:
            return ph_color
        name = self.clr_map.get(name, name)
        name = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}.get(name, name)
        return self.colors.get(name, "000000")

    def font(self, face):
        if not face:
            return None
        if face.startswith("+mj"):
            return self.major_ea if face.endswith("-ea") and self.major_ea else self.major
        if face.startswith("+mn"):
            return self.minor_ea if face.endswith("-ea") and self.minor_ea else self.minor
        return face


def resolve_color(el, theme, ph_color=None):
    """el: a colour choice element (srgbClr/schemeClr/...) or a parent holding one.
    Returns '#RRGGBB' or '#RRGGBBAA' or None."""
    if el is None:
        return None
    tag = etree.QName(el).localname
    if tag not in ("srgbClr", "schemeClr", "sysClr", "prstClr", "scrgbClr", "hslClr"):
        for c in el:
            if etree.QName(c).localname in ("srgbClr", "schemeClr", "sysClr", "prstClr", "scrgbClr"):
                return resolve_color(c, theme, ph_color)
        return None
    if tag == "srgbClr":
        base = el.get("val")
    elif tag == "schemeClr":
        base = theme.scheme(el.get("val"), ph_color.lstrip("#")[:6] if ph_color else None)
    elif tag == "sysClr":
        base = el.get("lastClr") or SYS_DEFAULT.get(el.get("val"), "000000")
    elif tag == "prstClr":
        base = PRESET_COLORS.get(el.get("val"), "000000")
    elif tag == "scrgbClr":
        base = "%02X%02X%02X" % tuple(int(int(el.get(k)) / 100000 * 255) for k in ("r", "g", "b"))
    else:
        base = "000000"
    r, g, b = _hex_to_rgb(base)
    alpha = 1.0
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    for m in el:
        mt = etree.QName(m).localname
        v = int(m.get("val", "0")) / 100000.0
        if mt == "lumMod":
            l = l * v
        elif mt == "lumOff":
            l = l + v
        elif mt == "tint":   # move toward white
            r, g, b = colorsys.hls_to_rgb(h, l, s)
            r, g, b = (c + (1 - c) * (1 - v) for c in (r, g, b))
            h, l, s = colorsys.rgb_to_hls(r, g, b)
        elif mt == "shade":  # move toward black
            r, g, b = colorsys.hls_to_rgb(h, l, s)
            r, g, b = (c * v for c in (r, g, b))
            h, l, s = colorsys.rgb_to_hls(r, g, b)
        elif mt == "satMod":
            s = s * v
        elif mt == "alpha":
            alpha = v
        l = max(0.0, min(1.0, l))
        s = max(0.0, min(1.0, s))
    out = _rgb_to_hex(colorsys.hls_to_rgb(h, l, s))
    if alpha < 0.999:
        out += "%02X" % round(alpha * 255)
    return out


def parse_fill(parent, theme, ph_color=None):
    """Look at direct children of parent (spPr / bgPr / tcPr) for a fill."""
    if parent is None:
        return None
    for c in parent:
        t = etree.QName(c).localname
        if t == "noFill":
            return {"type": "none"}
        if t == "grpFill":
            return {"type": "group"}
        if t == "solidFill":
            return {"type": "solid", "color": resolve_color(c, theme, ph_color)}
        if t == "gradFill":
            stops = []
            for gs in c.findall("a:gsLst/a:gs", NS):
                stops.append({"pos": int(gs.get("pos", 0)) / 1000.0,
                              "color": resolve_color(gs, theme, ph_color)})
            lin = c.find("a:lin", NS)
            path = c.find("a:path", NS)
            g = {"type": "gradient", "stops": sorted(stops, key=lambda s: s["pos"])}
            if path is not None:
                g["shape"] = "radial"
            else:
                g["shape"] = "linear"
                g["angle"] = int(lin.get("ang", 0)) / 60000.0 if lin is not None else 0
            return g
        if t == "blipFill":
            return {"type": "image", "rid": c.find("a:blip", NS).get(R + "embed") if c.find("a:blip", NS) is not None else None}
        if t == "pattFill":
            fg = resolve_color(c.find("a:fgClr", NS), theme, ph_color)
            return {"type": "solid", "color": fg}
    return None


def parse_line(spPr, theme, ph_color=None):
    if spPr is None:
        return None
    ln = spPr.find("a:ln", NS)
    if ln is None:
        return None
    out = {}
    if ln.get("w"):
        out["width"] = int(ln.get("w")) / EMU_PT
    f = parse_fill(ln, theme, ph_color)
    if f:
        out["fill"] = f
    dash = ln.find("a:prstDash", NS)
    if dash is not None:
        out["dash"] = dash.get("val")
    for end in ("headEnd", "tailEnd"):
        e = ln.find("a:" + end, NS)
        if e is not None and e.get("type", "none") != "none":
            out[end] = {"type": e.get("type"), "w": e.get("w", "med"), "len": e.get("len", "med")}
    return out


# ---------------------------------------------------------------- inherited text styles
def _lvl_props(container, lvl):
    if container is None:
        return None
    return container.find("a:lvl%dpPr" % (lvl + 1), NS)


class TextDefaults:
    """Resolves size/colour/font/bold/align for a run by walking the
    placeholder inheritance chain down to the master text styles."""

    def __init__(self, theme, master):
        self.theme = theme
        tx = master.element.find("p:txStyles", NS)
        self.styles = {
            "title": tx.find("p:titleStyle", NS) if tx is not None else None,
            "body": tx.find("p:bodyStyle", NS) if tx is not None else None,
            "other": tx.find("p:otherStyle", NS) if tx is not None else None,
        }

    def chain(self, shape_el, layout_ph_el, master_ph_el, ph_type):
        chain = []
        for el in (shape_el, layout_ph_el, master_ph_el):
            if el is not None:
                chain.append(el.find(".//p:txBody/a:lstStyle", NS))
        kind = "other"
        if ph_type in ("title", "ctrTitle"):
            kind = "title"
        elif ph_type is not None:
            kind = "body"
        chain.append(self.styles.get(kind))
        return chain


def first_attr(elements, attr):
    for e in elements:
        if e is not None and e.get(attr) is not None:
            return e.get(attr)
    return None


# ------------------------------------------------------------------------ extractor
class Extractor:
    def __init__(self, path, out_dir):
        self.prs = Presentation(path)
        self.out_dir = out_dir
        self.media_dir = os.path.join(out_dir, "media")
        os.makedirs(self.media_dir, exist_ok=True)
        self.sw = self.prs.slide_width / EMU_PT
        self.sh = self.prs.slide_height / EMU_PT
        self.media_count = 0

    # theme for a master
    def theme_for(self, master):
        theme_el = None
        for rel in master.part.rels.values():
            if rel.reltype.endswith("/theme"):
                theme_el = etree.fromstring(rel.target_part.blob)
        cm = master.element.find("p:clrMap", NS)
        clr_map = dict(cm.attrib) if cm is not None else {}
        return Theme(theme_el, clr_map)

    def save_blob(self, blob, ext):
        self.media_count += 1
        name = "img%03d.%s" % (self.media_count, ext)
        with open(os.path.join(self.media_dir, name), "wb") as f:
            f.write(blob)
        return "media/" + name

    # --------------------------------------------------------------- background
    def background(self, slide, theme):
        for src in (slide, slide.slide_layout, slide.slide_layout.slide_master):
            bg = src.element.find("p:cSld/p:bg", NS)
            if bg is None:
                continue
            bgPr = bg.find("p:bgPr", NS)
            if bgPr is not None:
                f = parse_fill(bgPr, theme)
                if f and f["type"] == "image" and f.get("rid"):
                    part = src.part.related_part(f["rid"])
                    f = {"type": "image", "src": self.save_blob(part.blob, part.partname.ext)}
                if f:
                    return f
            bgRef = bg.find("p:bgRef", NS)
            if bgRef is not None:
                color = resolve_color(bgRef, theme)
                idx = int(bgRef.get("idx", "1001"))
                styles = theme.bg_styles
                if idx >= 1001 and idx - 1001 < len(styles):
                    holder = etree.Element("x")
                    holder.append(copy.deepcopy(styles[idx - 1001]))
                    f = parse_fill(holder, theme, color)
                    if f and f["type"] != "image":
                        return f
                return {"type": "solid", "color": color}
        return {"type": "solid", "color": "#" + theme.scheme("bg1")}

    # --------------------------------------------------------------- animations
    def animations(self, slide):
        timing = slide.element.find("p:timing", NS)
        out = []
        if timing is None:
            return out
        main = None
        for ctn in timing.iter(P + "cTn"):
            if ctn.get("nodeType") == "mainSeq":
                main = ctn
                break
        if main is None:
            return out
        step = 0
        order = 0
        for click_par in main.findall("p:childTnLst/p:par", NS):
            step += 1
            for ctn in click_par.iter(P + "cTn"):
                pc = ctn.get("presetClass")
                if pc is None:
                    continue
                tgt = ctn.find(".//p:spTgt", NS)
                if tgt is None:
                    continue
                durs = [int(c.get("dur")) for c in ctn.iter(P + "cTn")
                        if c.get("dur") and c.get("dur").isdigit()]
                delay = 0
                cond = ctn.find("p:stCondLst/p:cond", NS)
                if cond is not None and cond.get("delay", "0").isdigit():
                    delay = int(cond.get("delay"))
                # the enclosing "with/after" group may carry the delay instead
                parent_ctn = ctn.getparent().getparent().getparent()
                if parent_ctn is not None and parent_ctn.tag == P + "cTn":
                    pcond = parent_ctn.find("p:stCondLst/p:cond", NS)
                    if pcond is not None and pcond.get("delay", "0").isdigit():
                        delay += int(pcond.get("delay"))
                pid = int(ctn.get("presetID", "0"))
                sub = int(ctn.get("presetSubtype", "0"))
                order += 1
                out.append({
                    "spid": tgt.get("spid"),
                    "class": pc,                       # entr / emph / exit / path
                    "presetID": pid,
                    "effect": ENTRANCE.get(pid, "fade") if pc == "entr" else ("pulse" if pc == "emph" else pc),
                    "direction": DIRECTION.get(sub, "bottom"),
                    "trigger": {"clickEffect": "click", "withEffect": "with",
                                "afterEffect": "after"}.get(ctn.get("nodeType"), "with"),
                    "step": step,
                    "order": order,
                    "delay": delay,
                    "dur": max(durs) if durs else 500,
                    "paragraph": tgt.find("p:txEl", NS) is not None,
                })
        return out

    # --------------------------------------------------------------- shapes
    def walk(self, shapes, xf, theme, tdefs, layout, origin, out):
        """xf: function (x,y,w,h in EMU) -> (x,y,w,h in pt) for group children."""
        for shp in shapes:
            el = shp._element
            tag = etree.QName(el).localname
            cnv = el.find(".//p:cNvPr", NS)
            if cnv is not None and cnv.get("hidden") in ("1", "true"):
                continue  # hidden in PowerPoint (think-cell markers, notes, old versions)
            if tag == "grpSp":
                gf = parse_fill(el.find("p:grpSpPr", NS), theme)
                outer = getattr(self, "_gfill", None)
                self._gfill = outer if (gf is None or gf["type"] == "group") else gf
                off = el.find("p:grpSpPr/a:xfrm/a:off", NS)
                ext = el.find("p:grpSpPr/a:xfrm/a:ext", NS)
                choff = el.find("p:grpSpPr/a:xfrm/a:chOff", NS)
                chext = el.find("p:grpSpPr/a:xfrm/a:chExt", NS)
                if None in (off, ext, choff, chext):
                    self.walk(shp.shapes, xf, theme, tdefs, layout, origin, out)
                    self._gfill = outer
                    continue
                ox, oy = int(off.get("x")), int(off.get("y"))
                ex, ey = int(ext.get("cx")), int(ext.get("cy"))
                cx, cy = int(choff.get("x")), int(choff.get("y"))
                cex, cey = max(1, int(chext.get("cx"))), max(1, int(chext.get("cy")))
                sx, sy = ex / cex, ey / cey

                def child_xf(x, y, w, h, _p=xf, ox=ox, oy=oy, cx=cx, cy=cy, sx=sx, sy=sy):
                    return _p(ox + (x - cx) * sx, oy + (y - cy) * sy, w * sx, h * sy)

                outer_grp = getattr(self, "_grp", None)
                self._grp = outer_grp or "g%s" % shp.shape_id  # top-level group id
                self.walk(shp.shapes, child_xf, theme, tdefs, layout, origin, out)
                self._grp = outer_grp
                self._gfill = outer
                continue
            try:
                item = self.element(shp, xf, theme, tdefs, layout, origin)
            except Exception as e:  # never let one odd shape kill the deck
                item = self.raster_item(shp, xf, origin, "error: %s" % e)
            if item:
                if getattr(self, "_grp", None):
                    item["group"] = self._grp
                if getattr(shp, "is_placeholder", False):
                    try:
                        t = shp._element.find(".//p:nvPr/p:ph", NS).get("type")
                        item["ph"] = t or "body"
                    except Exception:
                        pass
                out.append(item)

    def geom(self, shp):
        x, y, w, h = shp.left, shp.top, shp.width, shp.height
        if x is None or w is None:
            return None
        return x, y, w, h

    def raster_item(self, shp, xf, origin, reason):
        g = self.geom(shp)
        if g is None:
            return None
        x, y, w, h = xf(*g)
        return {"id": str(shp.shape_id), "name": shp.name, "kind": "raster", "reason": reason,
                "x": x, "y": y, "w": w, "h": h, "rot": 0, "origin": origin}

    def inherited_spPr(self, shp, layout):
        """spPr of the matching layout (then master) placeholder, for geometry/fill the slide leaves out."""
        if not getattr(shp, "is_placeholder", False):
            return []
        try:
            idx = shp.placeholder_format.idx
            ph_type = shp._element.find(".//p:nvPr/p:ph", NS).get("type")
        except Exception:
            return []
        out = []
        for src in (layout, layout.slide_master):
            for ph in src.placeholders:
                el = ph._element
                t = el.find(".//p:nvPr/p:ph", NS).get("type")
                if ph.placeholder_format.idx == idx or (ph_type and t == ph_type):
                    sp = el.find("p:spPr", NS)
                    if sp is not None:
                        out.append(sp)
                    break
        return out

    def geometry_el(self, spPr, inherited):
        for sp in [spPr] + inherited:
            if sp is None:
                continue
            g = sp.find("a:custGeom", NS)
            if g is None:
                g = sp.find("a:prstGeom", NS)
            if g is not None:
                return g
        return None

    def element(self, shp, xf, theme, tdefs, layout, origin):
        el = shp._element
        tag = etree.QName(el).localname
        g = self.geom(shp)
        if g is None:
            return None
        x, y, w, h = xf(*g)
        spPr = el.find("p:spPr", NS)
        xfrm = spPr.find("a:xfrm", NS) if spPr is not None else None
        rot = int(xfrm.get("rot", 0)) / 60000.0 if xfrm is not None else 0
        flipH = xfrm is not None and xfrm.get("flipH") == "1"
        flipV = xfrm is not None and xfrm.get("flipV") == "1"
        base = {"id": str(shp.shape_id), "name": shp.name, "x": x, "y": y, "w": w, "h": h,
                "rot": rot, "flipH": flipH, "flipV": flipV, "origin": origin}

        if tag == "graphicFrame":
            uri = el.find(".//a:graphicData", NS).get("uri", "")
            kind = "table" if uri.endswith("/table") else "chart" if "chart" in uri else "diagram" if "diagram" in uri else "object"
            return self.raster_item(shp, xf, origin, kind)
        if tag in ("contentPart",):
            return self.raster_item(shp, xf, origin, tag)

        style = el.find("p:style", NS)
        style_fill = style_line = style_font = None
        style_lw = None
        if style is not None:
            fr = style.find("a:fillRef", NS)
            lr = style.find("a:lnRef", NS)
            fo = style.find("a:fontRef", NS)
            if fr is not None and fr.get("idx") != "0":
                style_fill = {"type": "solid", "color": resolve_color(fr, theme)}
                idx = int(fr.get("idx"))
                if 1 <= idx <= len(theme.fill_styles):
                    holder = etree.Element("x")
                    holder.append(copy.deepcopy(theme.fill_styles[idx - 1]))
                    f = parse_fill(holder, theme, style_fill["color"])
                    if f and f["type"] in ("solid", "gradient"):
                        style_fill = f
            if lr is not None and lr.get("idx") != "0":
                style_line = resolve_color(lr, theme)
                idx = int(lr.get("idx"))
                style_lw = theme.ln_widths[idx - 1] / EMU_PT if 1 <= idx <= len(theme.ln_widths) else 0.75
            if fo is not None:
                style_font = resolve_color(fo, theme)

        if tag == "pic":
            blip = el.find(".//a:blip", NS)
            rid = blip.get(R + "embed") if blip is not None else None
            svg = blip.find(".//{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip", NS) if blip is not None else None
            if svg is not None and svg.get(R + "embed"):
                rid = svg.get(R + "embed")  # crisp vector version Office stores next to the PNG
            if not rid:
                return self.raster_item(shp, xf, origin, "linked picture")
            part = shp.part.related_part(rid)
            ext = part.partname.ext
            if ext.lower() in ("emf", "wmf", "tif", "tiff"):
                return self.raster_item(shp, xf, origin, "vector/tiff picture")
            src = self.save_blob(part.blob, ext)
            crop = {}
            sr = el.find(".//a:srcRect", NS)
            if sr is not None:
                for k in ("l", "t", "r", "b"):
                    crop[k] = int(sr.get(k, 0)) / 100000.0
            geo = self.geometry_el(spPr, self.inherited_spPr(shp, layout))
            is_cust = geo is not None and etree.QName(geo).localname == "custGeom"
            prst = geo if (geo is not None and not is_cust) else None
            item = dict(base, kind="picture", src=src, crop=crop,
                        prst=prst.get("prst") if prst is not None else "rect",
                        line=parse_line(spPr, theme))
            cust = geo if is_cust else None
            if cust is not None:  # photo cut into a letter / custom outline
                paths = self.custom_paths(cust)
                if paths is None:
                    return self.raster_item(shp, xf, origin, "picture with arc-shaped outline")
                item["prst"] = "custom"
                item["paths"] = paths
            elif item["prst"] not in ("rect", "ellipse", "roundRect", "flowChartConnector") and item["prst"] in SUPPORTED_PRST:
                item["adj"] = {}
            elif item["prst"] not in ("rect", "ellipse", "roundRect", "flowChartConnector"):
                return self.raster_item(shp, xf, origin, "picture with %s outline" % item["prst"])
            return item

        # sp / cxnSp
        inh = self.inherited_spPr(shp, layout)
        geo = self.geometry_el(spPr, inh)
        prst_el = geo if geo is not None and etree.QName(geo).localname == "prstGeom" else None
        cust = geo if geo is not None and etree.QName(geo).localname == "custGeom" else None
        prst = prst_el.get("prst") if prst_el is not None else None
        adj = {}
        if prst_el is not None:
            for gd in prst_el.findall("a:avLst/a:gd", NS):
                m = re.match(r"val\s+(-?\d+)", gd.get("fmla", ""))
                if m:
                    adj[gd.get("name")] = int(m.group(1))

        fill = parse_fill(spPr, theme)
        for sp in inh:
            if fill is not None:
                break
            fill = parse_fill(sp, theme)
        if fill and fill["type"] == "group":
            fill = getattr(self, "_gfill", None) or {"type": "none"}
        if fill and fill["type"] == "image" and fill.get("rid"):
            part = shp.part.related_part(fill["rid"])
            fill = {"type": "image", "src": self.save_blob(part.blob, part.partname.ext)}
        if fill is None:
            fill = style_fill or {"type": "none"}
        line = parse_line(spPr, theme) or {}
        if "fill" not in line:
            line["fill"] = {"type": "solid", "color": style_line} if style_line else {"type": "none"}
        if "width" not in line:
            line["width"] = style_lw or 0.75

        is_line = tag == "cxnSp" or prst in LINE_PRST
        if is_line:
            return dict(base, kind="line", prst=prst or "line", adj=adj, line=line)

        item = dict(base, kind="shape", prst=prst, adj=adj, fill=fill, line=line)
        if cust is not None:
            paths = self.custom_paths(cust)
            if paths is None:
                return self.raster_item(shp, xf, origin, "custom geometry with arcs")
            item["prst"] = "custom"
            item["paths"] = paths
        elif prst is None:
            item["prst"] = "rect"
        elif prst not in SUPPORTED_PRST:
            item["kind"] = "raster"
            item["reason"] = "geometry %s" % prst
        item["shadow"] = spPr is not None and spPr.find("a:effectLst/a:outerShdw", NS) is not None
        item["text"] = self.text(shp, theme, tdefs, layout, style_font)
        # callouts: geometry or a name like "말풍선"
        if prst in CALLOUT_PRST or CALLOUT_NAME_RE.search(shp.name or ""):
            item["callout"] = True
            if prst in ("wedgeRectCallout", "wedgeRoundRectCallout", "wedgeEllipseCallout", "cloudCallout"):
                a1 = adj.get("adj1", -20833)
                a2 = adj.get("adj2", 62500)
                tx = x + w / 2 + w * a1 / 100000.0
                ty = y + h / 2 + h * a2 / 100000.0
                if flipH:
                    tx = x + w / 2 - w * a1 / 100000.0
                if flipV:
                    ty = y + h / 2 - h * a2 / 100000.0
                item["tail"] = [tx, ty]
        if prst in ARROW_PRST:
            item["arrow"] = True
        return item

    def custom_paths(self, cust):
        paths = []
        for path in cust.findall("a:pathLst/a:path", NS):
            pw = int(path.get("w", 0)) or None
            ph = int(path.get("h", 0)) or None
            cmds = []
            for c in path:
                t = etree.QName(c).localname
                pts = [(int(pt.get("x")), int(pt.get("y"))) for pt in c.findall("a:pt", NS)]
                if t == "moveTo":
                    cmds.append(["M"] + list(pts[0]))
                elif t == "lnTo":
                    cmds.append(["L"] + list(pts[0]))
                elif t == "cubicBezTo":
                    cmds.append(["C"] + [v for p_ in pts for v in p_])
                elif t == "quadBezTo":
                    cmds.append(["Q"] + [v for p_ in pts for v in p_])
                elif t == "close":
                    cmds.append(["Z"])
                elif t == "arcTo":
                    return None
            paths.append({"w": pw, "h": ph, "fill": path.get("fill", "norm") != "none",
                          "stroke": path.get("stroke", "1") != "0", "d": cmds})
        return paths

    # --------------------------------------------------------------- text
    def text(self, shp, theme, tdefs, layout, style_font):
        if not shp.has_text_frame or not shp.text_frame.text.strip():
            return None
        el = shp._element
        txBody = el.find("p:txBody", NS)
        bodyPr = txBody.find("a:bodyPr", NS)
        ph_type = None
        layout_ph = master_ph = None
        if shp.is_placeholder:
            phf = shp.placeholder_format
            ph_type = etree.QName(el.find(".//p:nvPr/p:ph", NS)).localname and (
                el.find(".//p:nvPr/p:ph", NS).get("type") or "body")
            for lp in layout.placeholders:
                if lp.placeholder_format.idx == phf.idx:
                    layout_ph = lp._element
            master = layout.slide_master
            for mp in master.placeholders:
                mt = mp._element.find(".//p:nvPr/p:ph", NS).get("type") or "body"
                if mt == ph_type or (ph_type == "ctrTitle" and mt == "title") or (
                        ph_type == "subTitle" and mt == "body"):
                    master_ph = mp._element
        lst_chain = tdefs.chain(el, layout_ph, master_ph, ph_type)

        def body_attr(name, default):
            for src in (el, layout_ph, master_ph):
                if src is None:
                    continue
                bp = src.find(".//p:txBody/a:bodyPr", NS)
                if bp is not None and bp.get(name) is not None:
                    return bp.get(name)
            return default

        ins = {k[0]: int(body_attr(k, d)) / EMU_PT for k, d in
               (("lIns", 91440), ("tIns", 45720), ("rIns", 91440), ("bIns", 45720))}
        scale = 1.0
        na = bodyPr.find("a:normAutofit", NS) if bodyPr is not None else None
        if na is not None and na.get("fontScale"):
            scale = int(na.get("fontScale")) / 100000.0
        vert = body_attr("vert", "horz")
        out = {"anchor": body_attr("anchor", "t"), "inset": ins, "wrap": body_attr("wrap", "square") != "none",
               "vert": vert, "paragraphs": []}

        for p in txBody.findall("a:p", NS):
            pPr = p.find("a:pPr", NS)
            lvl = int(pPr.get("lvl", 0)) if pPr is not None else 0
            lvl_chain = [pPr] + [_lvl_props(c, lvl) for c in lst_chain]
            def_rprs = [e.find("a:defRPr", NS) if e is not None else None for e in lvl_chain]
            algn = first_attr(lvl_chain, "algn") or "l"
            ln_pct = None
            for e in lvl_chain:
                if e is not None:
                    sp = e.find("a:lnSpc/a:spcPct", NS)
                    if sp is not None:
                        ln_pct = int(sp.get("val")) / 100000.0
                        break
            spc_bef = None
            for e in lvl_chain:
                if e is not None:
                    sp = e.find("a:spcBef/a:spcPts", NS)
                    if sp is not None:
                        spc_bef = int(sp.get("val")) / 100.0
                        break
            bullet = None
            for e in lvl_chain:
                if e is None:
                    continue
                if e.find("a:buNone", NS) is not None:
                    break
                bc = e.find("a:buChar", NS)
                if bc is not None:
                    bullet = bc.get("char")
                    break
                if e.find("a:buAutoNum", NS) is not None:
                    bullet = "#"
                    break
            indent = first_attr(lvl_chain, "marL")
            hang = first_attr(lvl_chain, "indent")
            runs = []
            for r in p:
                t = etree.QName(r).localname
                if t == "br":
                    runs.append({"br": True})
                    continue
                if t not in ("r", "fld"):
                    continue
                rPr = r.find("a:rPr", NS)
                chain = [rPr] + def_rprs
                sz = first_attr(chain, "sz")
                color = None
                for e in chain:
                    if e is not None:
                        sf = e.find("a:solidFill", NS)
                        if sf is not None:
                            color = resolve_color(sf, theme)
                            break
                # Latin text uses the latin face, Hangul/CJK the ea face; the
                # browser does the same per glyph when both are in font-family
                latin = ea = None
                for e in chain:
                    if e is None:
                        continue
                    if latin is None:
                        f = e.find("a:latin", NS)
                        if f is not None and f.get("typeface"):
                            latin = theme.font(f.get("typeface"))
                    if ea is None:
                        f = e.find("a:ea", NS)
                        if f is not None and f.get("typeface"):
                            ea = theme.font(f.get("typeface"))
                is_title = ph_type in ("title", "ctrTitle")
                tnode = r.find("a:t", NS)
                text = tnode.text if tnode is not None and tnode.text else ""
                if t == "fld" and (r.get("type") or "").startswith("slidenum"):
                    text = "{{SLIDENUM}}"
                runs.append({
                    "text": text,
                    "size": (int(sz) / 100.0 if sz else 18.0) * scale,
                    "bold": first_attr(chain, "b") in ("1", "true"),
                    "italic": first_attr(chain, "i") in ("1", "true"),
                    "underline": (first_attr(chain, "u") or "none") != "none",
                    "color": color or style_font or "#" + theme.scheme("tx1"),
                    "font": latin or (theme.major if is_title else theme.minor),
                    "font_ea": ea or (theme.major_ea if is_title else theme.minor_ea),
                })
            end = p.find("a:endParaRPr", NS)
            empty_size = None
            if not runs and end is not None and end.get("sz"):
                empty_size = int(end.get("sz")) / 100.0 * scale
            out["paragraphs"].append({"align": algn, "line": ln_pct, "spaceBefore": spc_bef,
                                      "bullet": bullet, "level": lvl,
                                      "marL": int(indent) / EMU_PT if indent else 0,
                                      "indent": int(hang) / EMU_PT if hang else 0,
                                      "runs": runs, "emptySize": empty_size})
        return out

    # --------------------------------------------------------------- deck
    def run(self):
        slides = []
        for i, slide in enumerate(self.prs.slides, 1):
            layout = slide.slide_layout
            master = layout.slide_master
            theme = self.theme_for(master)
            tdefs = TextDefaults(theme, master)
            ident = lambda x, y, w, h: (x / EMU_PT, y / EMU_PT, w / EMU_PT, h / EMU_PT)
            elements = []
            show_master = slide.element.find("p:cSld", NS).getparent().get("showMasterSp", "1") != "0"
            layout_show = layout.element.get("showMasterSp", "1") != "0"
            if show_master:
                if layout_show:
                    self.walk([s for s in master.shapes if not s.is_placeholder], ident, theme, tdefs, layout, "master", elements)
                self.walk([s for s in layout.shapes if not s.is_placeholder], ident, theme, tdefs, layout, "layout", elements)
            self.walk(slide.shapes, ident, theme, tdefs, layout, "slide", elements)
            notes = ""
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else ""
            for e in elements:
                for para in (e.get("text") or {}).get("paragraphs", []):
                    for r in para["runs"]:
                        if r.get("text") == "{{SLIDENUM}}":
                            r["text"] = str(i)
            slides.append({
                "index": i,
                "background": self.background(slide, theme),
                "elements": elements,
                "animations": self.animations(slide),
                "notes": notes,
                "hidden": slide.element.get("show") == "0",
            })
        data = {"width": self.sw, "height": self.sh, "slides": slides}
        with open(os.path.join(self.out_dir, "slides.json"), "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=1)
        return data


def summary(data):
    lines = ["slide size: %.0f x %.0f pt" % (data["width"], data["height"])]
    for s in data["slides"]:
        kinds = {}
        for e in s["elements"]:
            kinds[e["kind"]] = kinds.get(e["kind"], 0) + 1
        callouts = [e["name"] for e in s["elements"] if e.get("callout")]
        arrows = sum(1 for e in s["elements"] if e["kind"] == "line" and (
            e["line"].get("headEnd") or e["line"].get("tailEnd"))) + sum(
            1 for e in s["elements"] if e.get("arrow"))
        raster = ["%s (%s)" % (e["name"], e.get("reason")) for e in s["elements"] if e["kind"] == "raster"]
        lines.append("slide %d%s: %s | callouts %d %s | arrows %d | animations %d%s" % (
            s["index"], " [hidden]" if s["hidden"] else "", kinds, len(callouts), callouts, arrows,
            len(s["animations"]), (" | raster: " + ", ".join(raster)) if raster else ""))
    return "\n".join(lines)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    os.makedirs(sys.argv[2], exist_ok=True)
    d = Extractor(sys.argv[1], sys.argv[2]).run()
    print(summary(d))
